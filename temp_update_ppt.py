from pptx import Presentation
from pptx.util import Inches
from pathlib import Path

ppt_path = Path(r'C:\application\AI_Project\SentinelAI\SentinelAI_Atos_Presentation - base.pptx')
prs = Presentation(ppt_path)


def update_slide_text(slide, title, body):
    text_shapes = [s for s in slide.shapes if getattr(s, 'has_text_frame', False) and s.has_text_frame and s.text.strip()]
    main_shapes = []
    for s in text_shapes:
        t = s.text.strip()
        if 'Atos Confidential' in t or 'May 2026' in t or (t.startswith('SentinelAI') and '|' in t):
            continue
        main_shapes.append(s)

    if not main_shapes:
        return

    main_shapes[0].text = title
    if len(main_shapes) > 1:
        main_shapes[1].text = body
    else:
        # append body below title if only one main text box exists
        tf = main_shapes[0].text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = title
        p.level = 0
        p2 = tf.add_paragraph()
        p2.text = body
        p2.level = 0


update_slide_text(prs.slides[1], 'Why an AI Layer on Top of Observability?', 'AIOps theory:\n• Observability data is raw signal, not insight\n• SentinelAI adds memory, reasoning and live investigation\n• The platform turns logs into RCA, context and action\n• Knowledge Service makes each incident reusable for the next one')

update_slide_text(prs.slides[3], 'SentinelAI in One Slide', 'An AI-assisted incident investigation platform built on the stack we already use\n\n• Log RCA: paste or upload logs, generate issue / root cause / impacted service / fix\n• ELK live investigation: query live logs with service, severity and timeframe\n• Knowledge Service: adds engineering context from deployments, releases, timeline and Jira\n• Future: agentic investigation with context-aware reasoning and workflow actions')

update_slide_text(prs.slides[4], 'High-Level Architecture', 'Four deployable services + shared infrastructure\n\n• sentinelai-ui (React): RCA and ELK investigation views\n• SentinelAI-Core (Spring Boot): orchestration, RCA, ELK integration and SKS calls\n• SentinelAI-Engine (FastAPI): provider routing for Groq, Ollama, OpenAI and HuggingFace\n• SentinelAI-Knowledge-Service (Spring Boot): timeline, releases, deployments and Jira knowledge\n\nInfrastructure: PostgreSQL + pgvector, Elasticsearch, Ollama, Docker\n\nDraw.io reference: SentinelAI-Architecture-v2.drawio')

# add link box on slide 5
slide = prs.slides[4]
box = slide.shapes.add_textbox(Inches(0.55), Inches(7.0), Inches(4.8), Inches(0.35))
tf = box.text_frame
tf.clear()
p = tf.paragraphs[0]
r = p.add_run()
r.text = 'Open architecture draw.io'
r.hyperlink.address = r'C:\application\AI_Project\SentinelAI\SentinelAI-Architecture-v2.drawio'
p.alignment = 0

update_slide_text(prs.slides[7], 'Existing Capability – Engineering Context & Knowledge', 'Every investigation becomes reusable knowledge\n\n• Core can call SentinelAI-Knowledge-Service for timeline, deployment, release and Jira evidence\n• That context is appended to the RCA prompt for richer analysis\n• The UI displays engineering evidence alongside the RCA result\n• Over time, SentinelAI becomes a living memory layer for incident response')

update_slide_text(prs.slides[12], 'Roadmap — P1 Enhancements', 'Three capabilities that turn SentinelAI into a core engineering tool\n\n• Natural language operational search over ELK and engineering knowledge\n• Incident correlation using ELK logs, deployments, releases and timeline events from SKS\n• Deployment regression detection with richer context from changes and incidents')

update_slide_text(prs.slides[13], 'Roadmap — P2 Enhancements', 'Real-time AI monitoring · ITSM automation · Enterprise readiness\n\n• Real-time ELK/Kafka monitoring with AI-driven investigation\n• Knowledge Service becomes the enterprise memory backbone for incident context\n• Jira / Slack / Teams automation with RCA and runbook draft generation\n• RBAC, audit logging and LLMOps for production rollout')

update_slide_text(prs.slides[14], 'Future Vision — Enterprise AI SRE Copilot', 'SentinelAI as the central AI brain for production operations\n\n• The agent will reason over ELK, Kafka, deployments, releases and engineering knowledge\n• SentinelAI-Knowledge-Service becomes the memory layer for past incidents and change history\n• The human approves in minutes while the agent drafts RCA, runbooks and Jira actions\n• This moves the platform from demo to an enterprise-ready operations copilot')

prs.save(ppt_path)
print('Updated PPTX successfully:', ppt_path)
