from pptx import Presentation
from pptx.util import Inches
from pptx.enum.text import PP_ALIGN
from pathlib import Path

ppt_path = Path(r'C:\application\AI_Project\SentinelAI\SentinelAI_Atos_Presentation - base.pptx')
prs = Presentation(ppt_path)

# Helper to replace the first two non-footer text boxes on a slide

def replace_slide_body(slide, title, body):
    text_boxes = []
    for shape in slide.shapes:
        if getattr(shape, 'has_text_frame', False) and shape.has_text_frame and shape.text.strip():
            text_boxes.append(shape)

    # Prefer the first two text boxes that are not footer metadata
    target_boxes = []
    for shape in text_boxes:
        t = shape.text.strip()
        if 'Atos Confidential' in t or 'May 2026' in t or (t.startswith('SentinelAI') and '|' in t):
            continue
        target_boxes.append(shape)

    if not target_boxes:
        return

    target_boxes[0].text = title
    if len(target_boxes) > 1:
        target_boxes[1].text = body
    else:
        # append to first box if only one target box exists
        tf = target_boxes[0].text_frame
        tf.clear()
        p = tf.paragraphs[0]
        p.text = title
        p2 = tf.add_paragraph()
        p2.text = body


replace_slide_body(prs.slides[1], 'Why an AI Layer on Top of Observability?', 'AIOps theory:\n• Observability data is raw signal, not insight\n• SentinelAI adds memory, reasoning and live investigation\n• The platform turns logs into RCA, context and action\n• Knowledge Service makes each incident reusable for the next one')

replace_slide_body(prs.slides[3], 'SentinelAI in One Slide', 'An AI-assisted incident investigation platform built on the stack we already use\n\n• Log RCA: paste or upload logs, generate issue / root cause / impacted service / fix\n• ELK live investigation: query live logs with service, severity and timeframe\n• Knowledge Service: adds engineering context from deployments, releases, timeline and Jira\n• Future: agentic investigation with context-aware reasoning and workflow actions')

replace_slide_body(prs.slides[4], 'High-Level Architecture', 'Four deployable services + shared infrastructure\n\n• sentinelai-ui (React): RCA and ELK investigation views\n• SentinelAI-Core (Spring Boot): orchestration, RCA, ELK integration and SKS calls\n• SentinelAI-Engine (FastAPI): provider routing for Groq, Ollama, OpenAI and HuggingFace\n• SentinelAI-Knowledge-Service (Spring Boot): timeline, releases, deployments and Jira knowledge\n\nInfrastructure: PostgreSQL + pgvector, Elasticsearch, Ollama, Docker\n\nDraw.io reference: SentinelAI-Architecture-v2.drawio')

# add a small hyperlink textbox to the architecture slide
slide = prs.slides[4]
box = slide.shapes.add_textbox(Inches(0.55), Inches(7.0), Inches(4.6), Inches(0.35))
tf = box.text_frame
tf.clear()
p = tf.paragraphs[0]
r = p.add_run()
r.text = 'Open architecture draw.io'
r.hyperlink.address = r'C:\application\AI_Project\SentinelAI\SentinelAI-Architecture-v2.drawio'
p.alignment = PP_ALIGN.LEFT

replace_slide_body(prs.slides[7], 'Existing Capability – Engineering Context & Knowledge', 'Every investigation becomes reusable knowledge\n\n• Core can call SentinelAI-Knowledge-Service for timeline, deployment, release and Jira evidence\n• That context is appended to the RCA prompt for richer analysis\n• The UI displays engineering evidence alongside the RCA result\n• Over time, SentinelAI becomes a living memory layer for incident response')

replace_slide_body(prs.slides[12], 'Roadmap — P1 Enhancements', 'Three capabilities that turn SentinelAI into a core engineering tool\n\n• Natural language operational search over ELK and engineering knowledge\n• Incident correlation using ELK logs, deployments, releases and timeline events from SKS\n• Deployment regression detection with richer context from changes and incidents')

replace_slide_body(prs.slides[13], 'Roadmap — P2 Enhancements', 'Real-time AI monitoring · ITSM automation · Enterprise readiness\n\n• Real-time ELK/Kafka monitoring with AI-driven investigation\n• Knowledge Service becomes the enterprise memory backbone for incident context\n• Jira / Slack / Teams automation with RCA and runbook draft generation\n• RBAC, audit logging and LLMOps for production rollout')

replace_slide_body(prs.slides[14], 'Future Vision — Enterprise AI SRE Copilot', 'SentinelAI as the central AI brain for production operations\n\n• The agent will reason over ELK, Kafka, deployments, releases and engineering knowledge\n• SentinelAI-Knowledge-Service becomes the memory layer for past incidents and change history\n• The human approves in minutes while the agent drafts RCA, runbooks and Jira actions\n• This moves the platform from demo to an enterprise-ready operations copilot')

prs.save(ppt_path)
print('Updated PPTX successfully:', ppt_path)
