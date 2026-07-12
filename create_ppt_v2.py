from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_AUTO_SIZE
from pptx.dml.color import RGBColor
from pathlib import Path

out_path = Path(r'C:\application\AI_Project\SentinelAI\SentinelAI_Atos_Presentation - v2.pptx')
prs = Presentation()

# Theme colors
NAVY = RGBColor(21, 41, 74)
BLUE = RGBColor(0, 102, 204)
GRAY = RGBColor(90, 100, 110)
LIGHT = RGBColor(245, 248, 252)


def set_title(slide, title, subtitle=None):
    title_box = slide.shapes.title
    if title_box is None:
        title_box = slide.shapes.add_textbox(Inches(0.6), Inches(0.3), Inches(11.2), Inches(0.8))
    tf = title_box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.alignment = PP_ALIGN.LEFT
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(11)
        p2.font.color.rgb = GRAY
        p2.alignment = PP_ALIGN.LEFT


def add_bullets(slide, bullets, left=Inches(0.9), top=Inches(1.35), width=Inches(10.8), height=Inches(5.8), font_size=18):
    text_box = slide.shapes.add_textbox(left, top, width, height)
    tf = text_box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    for i, bullet in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = bullet
        p.level = 0
        p.font.size = Pt(font_size)
        p.font.color.rgb = RGBColor(35, 35, 35)
        p.font.name = 'Calibri'
        p.space_after = Pt(6)
        p.alignment = PP_ALIGN.LEFT
        p.bullet = True
    return text_box


def add_footer(slide, text):
    box = slide.shapes.add_textbox(Inches(0.6), Inches(7.1), Inches(10.8), Inches(0.35))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(10)
    p.font.color.rgb = GRAY
    p.alignment = PP_ALIGN.RIGHT


def add_link_box(slide, text, address, left=Inches(8.8), top=Inches(7.05), width=Inches(2.5), height=Inches(0.35)):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.hyperlink.address = address
    p.font.size = Pt(10)
    p.font.color.rgb = BLUE
    p.alignment = PP_ALIGN.RIGHT

# Slide 1
slide = prs.slides.add_slide(prs.slide_layouts[0])
slide.shapes.title.text = 'SentinelAI'
slide.placeholders[1].text = 'AI-Assisted Incident Investigation & RCA\nBuilt on ELK, Kafka, Cloud-Native Platforms, and Engineering Knowledge'
# subtitle in footer
add_footer(slide, 'Praveen Kumar Deshmukh | Atos | May 2026')

# Slide 2
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Why an AI Layer on Top of Observability?', 'AIOps theory and the business case')
add_bullets(slide, [
    'Observability data is raw signal, not insight. The value comes when we convert logs, traces and events into decisions.',
    'SentinelAI adds three layers: memory, reasoning and live investigation. That is the core difference from classic dashboards.',
    'The platform turns noisy telemetry into structured RCA: issue, root cause, impacted service and recommended fix.',
    'With SentinelAI-Knowledge-Service, each investigation becomes reusable knowledge for the next one.',
    'This helps reduce MTTR, improve consistency and reduce dependence on tribal knowledge.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 3
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Current Pain Points in Incident Investigation', 'Why the market and our teams need this capability')
add_bullets(slide, [
    'Engineers manually query ELK, correlate timestamps, and reconstruct incidents across many tabs.',
    'Similar incidents are investigated from scratch every time because there is no durable memory layer.',
    'Junior engineers lack context while senior engineers carry tribal knowledge in their heads.',
    'RCAs are often written late, inconsistently, and with incomplete evidence.',
    'The opportunity: turn operational data into a reusable, explainable investigation workflow.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 4
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'SentinelAI in One Slide', 'What the platform does today')
add_bullets(slide, [
    'Log RCA: paste or upload logs and generate issue, root cause, impacted service and suggested fix.',
    'ELK live investigation: query live logs using service, severity and timeframe without manual Kibana DSL.',
    'Engineering context enrichment: use deployment, release, timeline and Jira evidence to strengthen the RCA.',
    'Knowledge Service: collect and normalize engineering metadata from GitHub, Jira, Confluence and Jenkins.',
    'Future direction: move from assisted investigation to agentic operations with tool use and workflow actions.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 5
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'High-Level Architecture', 'Four deployable services with shared infrastructure')
add_bullets(slide, [
    'sentinelai-ui (React): Log RCA and ELK Investigation experiences with evidence-rich output.',
    'SentinelAI-Core (Spring Boot): orchestrates RCA, log ingestion, ELK integration and Knowledge Service calls.',
    'SentinelAI-Engine (FastAPI): routes analysis to Groq, Ollama, OpenAI or HuggingFace providers.',
    'SentinelAI-Knowledge-Service (Spring Boot): builds the engineering memory layer from timelines, deployments, releases and Jira.',
    'Infrastructure: PostgreSQL + pgvector, Elasticsearch, Ollama and Docker-based local deployment.'
])
add_link_box(slide, 'Open architecture draw.io', str(Path(r'C:\application\AI_Project\SentinelAI\SentinelAI-Architecture-v2.drawio').resolve()))
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 6 new knowledge service slide
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'SentinelAI-Knowledge-Service', 'The new capability added after the initial RCA MVP')
add_bullets(slide, [
    'Purpose: create a normalized engineering knowledge index for incident investigation and RCA enrichment.',
    'Capabilities: sync metadata from GitHub, Jira, Confluence and Jenkins; ingest webhooks; expose timeline and relationship APIs.',
    'Business value: provide deployment, release and issue context to the RCA flow without requiring the engineer to search across tools manually.',
    'Integration point: SentinelAI-Core can call the Knowledge Service to enrich prompts with recent timeline events, deployments, releases and Jira issues.',
    'Outcome: the system becomes a living memory layer for operations, not just a single-prompt LLM experience.'
])
add_footer(slide, 'New capability | SKS | May 2026')

# Slide 7
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Existing Capability – Log RCA', 'Paste logs and receive structured root-cause analysis')
add_bullets(slide, [
    'Engineer pastes log text or uploads a file from the UI.',
    'The UI sends the content to Core via POST /api/rca/analyze.',
    'Core optionally adds historical context from prior incidents and engineering evidence.',
    'The AI Engine returns structured fields such as issue, root cause, impacted service and recommended fix.',
    'This accelerates first-pass triage and gives a consistent RCA structure for downstream collaboration.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 8
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Existing Capability – ELK-Based Live Investigation', 'Investigate live logs directly from observability data')
add_bullets(slide, [
    'The engineer selects service, severity and timeframe.',
    'Core queries Elasticsearch and retrieves recent log lines for the target context.',
    'The AI Engine summarizes the findings and highlights suspicious log evidence.',
    'The output includes probable root cause, impacted service and recommended action.',
    'This removes the need to manually craft Kibana DSL queries for common investigations.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 9
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Existing Capability – RAG and Incident Memory', 'Every investigation becomes better for the next one')
add_bullets(slide, [
    'Uploaded incidents and logs can be stored together with embeddings for semantic retrieval.',
    'The Core service can search stored incidents for similar patterns and add them as context.',
    'This improves grounding and reduces hallucination risk for RCA generation.',
    'The Knowledge Service extends this further by adding deployment, release and incident timeline context.',
    'The result is a growing knowledge base that compounds value over time.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 10
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Pluggable AI Provider Architecture', 'Choice of provider without changing business logic')
add_bullets(slide, [
    'Groq offers low-latency cloud inference for fast interactions and demos.',
    'Ollama supports fully local execution for data residency and privacy-sensitive environments.',
    'OpenAI and HuggingFace provide enterprise and research-oriented alternatives.',
    'Core remains provider-agnostic; switching models is a configuration change, not a code rewrite.',
    'This gives flexibility for PoCs, pilot environments and eventual enterprise rollout.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 11
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Agentic AI – The Next Step', 'From assisted investigation to an SRE copilot')
add_bullets(slide, [
    'The next phase is a reasoning agent that plans multiple steps instead of making one-shot predictions.',
    'It can query ELK, inspect deployment history, search the incident knowledge base and draft a runbook or Jira action.',
    'The agent can loop over tool results, observe outcomes and refine its plan before responding.',
    'This is the move from "generate answer" to "investigate and act".',
    'The Knowledge Service becomes the memory backbone for that agentic workflow.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 12
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Roadmap – P1 Enhancements', 'What should come next in the first real product phase')
add_bullets(slide, [
    'Natural language operational search over ELK and engineering knowledge.',
    'Incident correlation across ELK, deployment, release and timeline events.',
    'Deployment regression detection that compares changes and signal patterns before and after release.',
    'Improved evidence ranking so the RCA output is both explainable and grounded.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 13
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Roadmap – P2 Enhancements', 'From pilot to product-grade operations platform')
add_bullets(slide, [
    'Real-time ELK and Kafka monitoring with AI-driven triage and anomaly grouping.',
    'Automation into Jira, Slack or Teams with RCA, runbook and ticket draft generation.',
    'Enterprise readiness with RBAC, audit logging, SSO and multi-tenant configuration.',
    'LLMOps and feedback loops so model quality, latency and cost can be managed over time.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 14
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Future Vision – Enterprise AI SRE Copilot', 'The long-term operating model')
add_bullets(slide, [
    'The on-call engineer receives a proactive AI investigation assistant instead of a blank dashboard.',
    'The agent reasons over ELK, Kafka, deployments, releases and incident history to propose next steps.',
    'SentinelAI-Knowledge-Service becomes the memory layer that keeps previous incidents and changes reusable.',
    'The human approves in minutes while the system drafts RCA, runbooks and workflow actions.',
    'This shifts the platform from a demo to a reusable enterprise operations accelerator.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 15
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Next Steps and Ask', 'How to turn this into a real Atos engagement')
add_bullets(slide, [
    'Choose one pilot environment and one incident workflow to prove the value in practice.',
    'Prioritize the first product features: operational search, incident correlation and deployment regression detection.',
    'Support the effort with access to real ELK, Kafka and deployment data for validation.',
    'Use this as a foundation for a reusable internal accelerator and client-facing offer.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

# Slide 16
slide = prs.slides.add_slide(prs.slide_layouts[1])
set_title(slide, 'Thank You', 'Questions and discussion')
add_bullets(slide, [
    'SentinelAI combines observability, AI reasoning and engineering knowledge into one investigation workflow.',
    'The new Knowledge Service makes the platform more grounded, more reusable and closer to an enterprise SRE copilot.',
    'The goal is not only to generate RCA, but to build a durable operational memory layer for the future.'
])
add_footer(slide, 'SentinelAI | Atos Confidential | May 2026')

prs.save(out_path)
print('Created', out_path)
print('Slide count', len(prs.slides))
