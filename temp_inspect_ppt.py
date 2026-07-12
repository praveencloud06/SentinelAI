from pptx import Presentation
from pathlib import Path
p = Path(r'C:\application\AI_Project\SentinelAI\SentinelAI_Atos_Presentation - base.pptx')
prs = Presentation(p)
print('slides', len(prs.slides))
for i, slide in enumerate(prs.slides, 1):
    print(f'--- slide {i} ---')
    print('title:', slide.shapes.title.text if slide.shapes.title else 'NO TITLE')
    for shape in slide.shapes:
        if hasattr(shape, 'text') and shape.text.strip():
            text = shape.text.replace('\n', ' | ').strip()
            print(text)
    print()
