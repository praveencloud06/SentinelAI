#!/usr/bin/env python3
"""
SentinelAI – Atos Internal Presentation Generator  v2
──────────────────────────────────────────────────────
10 lean slides with image placeholders for draw.io diagrams.
Requirements:  pip install python-pptx
Run:           python generate_ppt.py
Output:        SentinelAI_Atos_Presentation.pptx

HOW TO ADD DIAGRAMS (do this in PowerPoint / Google Slides after opening the file):
  1. Open the relevant .drawio file in draw.io / diagrams.net
  2. File → Export As → PNG  (Scale: 2x, Transparent: off)
  3. In PowerPoint: click the placeholder box → Delete → Insert → Pictures
     In Google Slides: click placeholder box → Delete → Insert → Image
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ── Colour palette ─────────────────────────────────────────────────────────
NAVY   = RGBColor(0x0D, 0x1B, 0x2A)
BLUE   = RGBColor(0x00, 0x6E, 0xD6)
CYAN   = RGBColor(0x00, 0xBC, 0xD4)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
LGRAY  = RGBColor(0xB0, 0xBB, 0xC9)
YELLOW = RGBColor(0xFF, 0xC1, 0x07)
GREEN  = RGBColor(0x4C, 0xAF, 0x50)
ORANGE = RGBColor(0xFF, 0x6D, 0x00)
PURPLE = RGBColor(0x7C, 0x4D, 0xFF)
DKBLUE = RGBColor(0x10, 0x24, 0x3E)
PLHOLD = RGBColor(0x06, 0x14, 0x26)   # image-placeholder background

W = Inches(13.33)
H = Inches(7.5)

# ── Layout constants ───────────────────────────────────────────────────────
CONTENT_TOP = Inches(1.38)
CONTENT_BOT = Inches(7.1)
CONTENT_H   = CONTENT_BOT - CONTENT_TOP   # ~5.72"

# Split-slide columns
L_LEFT  = Inches(0.4)
W_LEFT  = Inches(5.2)
L_RIGHT = Inches(5.85)
W_RIGHT = Inches(7.1)


# ── Low-level helpers ──────────────────────────────────────────────────────

def blank_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def rect(slide, l, t, w, h, fill=None, line=None, lw=Pt(1.5)):
    shp = slide.shapes.add_shape(1, l, t, w, h)
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = line
        shp.line.width = lw
    else:
        shp.line.fill.background()
    return shp


def tb(slide, text, l, t, w, h,
       size=14, bold=False, italic=False,
       color=WHITE, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf  = box.text_frame
    tf.word_wrap = True
    p   = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size   = Pt(size)
    r.font.bold   = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def hdr(slide, title, sub=None, accent=BLUE):
    """Standard top-bar + title used on every content slide."""
    rect(slide, 0, 0, W, Inches(0.07), fill=accent)
    tb(slide, title,
       Inches(0.5), Inches(0.12), Inches(12.3), Inches(0.65),
       size=25, bold=True, color=WHITE)
    if sub:
        tb(slide, sub,
           Inches(0.5), Inches(0.78), Inches(12.3), Inches(0.38),
           size=12, italic=True, color=CYAN)
    rect(slide, 0, H - Inches(0.07), W, Inches(0.07), fill=accent)
    tb(slide, "SentinelAI  |  Atos Confidential  |  May 2026",
       Inches(8.0), H - Inches(0.36), Inches(5.0), Inches(0.3),
       size=9, color=LGRAY, align=PP_ALIGN.RIGHT)


def img_placeholder(slide, l, t, w, h, filename):
    """
    Labelled drop zone for a draw.io exported PNG.
    Replace this box in PowerPoint / Google Slides with the exported image.
    """
    shp = slide.shapes.add_shape(1, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = PLHOLD
    shp.line.color.rgb = RGBColor(0x40, 0x55, 0x70)
    shp.line.width = Pt(1.5)

    mid_t = t + h / 2
    tb(slide, "INSERT DIAGRAM HERE",
       l + Inches(0.3), mid_t - Inches(0.75), w - Inches(0.6), Inches(0.48),
       size=15, bold=True, color=RGBColor(0x45, 0x58, 0x68), align=PP_ALIGN.CENTER)
    tb(slide, filename,
       l + Inches(0.2), mid_t - Inches(0.18), w - Inches(0.4), Inches(0.44),
       size=12, bold=True, italic=True, color=CYAN, align=PP_ALIGN.CENTER)
    tb(slide, "Export as PNG from draw.io  \u2192  delete this box  \u2192  insert picture",
       l + Inches(0.2), mid_t + Inches(0.32), w - Inches(0.4), Inches(0.34),
       size=10, color=LGRAY, align=PP_ALIGN.CENTER)


def blist(slide, items, l, t, w, size=13, color=WHITE, gap=Inches(0.58)):
    """Render bullet list; '  ' prefix = sub-bullet (smaller, gray)."""
    for i, item in enumerate(items):
        sub = item.startswith("  ")
        txt = ("    \u25e6  " + item.strip()) if sub else ("\u2022  " + item)
        tb(slide, txt, l, t + i * gap, w, gap,
           size=(size - 1) if sub else size,
           color=LGRAY if sub else color)


# ── Slide builders ─────────────────────────────────────────────────────────

def s01_title(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    rect(s, 0, 0, Inches(0.12), H, fill=BLUE)
    rect(s, 0, H - Inches(0.1), W, Inches(0.1), fill=BLUE)

    tb(s, "SentinelAI",
       Inches(0.6), Inches(1.0), Inches(8.5), Inches(1.5),
       size=58, bold=True, color=WHITE)
    tb(s, "AI-Assisted Incident Investigation & RCA",
       Inches(0.6), Inches(2.55), Inches(8.5), Inches(0.7),
       size=22, color=CYAN)
    rect(s, Inches(0.6), Inches(3.35), Inches(6.5), Inches(0.07), fill=CYAN)
    tb(s, "AI layer on top of ELK  \u00b7  Kafka  \u00b7  Cloud-Native Platforms",
       Inches(0.6), Inches(3.52), Inches(8.5), Inches(0.48),
       size=15, italic=True, color=LGRAY)

    tb(s, "Praveen Kumar Deshmukh",
       Inches(0.6), Inches(5.1), Inches(8.0), Inches(0.5),
       size=18, bold=True, color=WHITE)
    tb(s, "Principal Software Engineer & Architect  |  Atos",
       Inches(0.6), Inches(5.62), Inches(8.0), Inches(0.4),
       size=13, color=LGRAY)
    tb(s, "May 2026",
       Inches(0.6), Inches(6.1), Inches(3.0), Inches(0.38),
       size=13, color=CYAN)

    # Decorative stripes on the right
    for i, (ox, col) in enumerate([
        (9.0, CYAN), (9.6, BLUE), (9.3, CYAN), (9.8, BLUE),
        (9.1, CYAN), (9.5, BLUE), (9.0, CYAN), (9.7, BLUE),
    ]):
        rect(s, Inches(ox), Inches(1.2 + i * 0.75), Inches(13.0 - ox), Inches(0.06), fill=col)

    tb(s, "[ ATOS INTERNAL \u2014 CONFIDENTIAL ]",
       Inches(9.0), Inches(6.5), Inches(4.1), Inches(0.4),
       size=9, color=LGRAY, align=PP_ALIGN.CENTER)


def s02_problem_why(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "Problem & Why Now",
        "Engineers spend 40\u201360 % of incident response manually querying ELK and writing RCAs")

    # Thin vertical separator
    rect(s, Inches(5.67), CONTENT_TOP, Inches(0.04), Inches(5.3),
         fill=RGBColor(0x1E, 0x2E, 0x42))

    # Left col
    tb(s, "Current Pain Points",
       L_LEFT, CONTENT_TOP, W_LEFT, Inches(0.4),
       size=14, bold=True, color=ORANGE)
    pains = [
        "Manual log analysis — slow, repetitive, error-prone",
        "Kibana DSL expertise required; non-experts blocked",
        "RCAs written hours later from fading memory",
        "Past incidents not reused — same root causes recur",
        "Senior engineers bottleneck every on-call rotation",
    ]
    blist(s, pains, L_LEFT, CONTENT_TOP + Inches(0.5), W_LEFT,
          size=13, gap=Inches(0.78))

    # Right col
    tb(s, "Why Act Now",
       L_RIGHT, CONTENT_TOP, W_RIGHT, Inches(0.4),
       size=14, bold=True, color=CYAN)
    whys = [
        "Datadog AI, Dynatrace Davis, New Relic Grok — vendors racing to ship AI copilots",
        "Early adopters report 30\u201350 % MTTR reduction with AI-assisted triage",
        "Our ELK + Kafka + Spring Boot stack is already AI-ready — no infra changes",
        "SentinelAI is a working system, not a prototype — ready to grow",
        "Window to differentiate Atos proposals is open right now",
    ]
    blist(s, whys, L_RIGHT, CONTENT_TOP + Inches(0.5), W_RIGHT,
          size=13, gap=Inches(0.78))

    rect(s, 0, H - Inches(0.88), W, Inches(0.72), fill=RGBColor(0x00, 0x38, 0x60))
    tb(s, "The question is not whether to add AI to observability \u2014 it is who builds it first.",
       Inches(0.5), H - Inches(0.84), Inches(12.3), Inches(0.6),
       size=14, bold=True, color=YELLOW, align=PP_ALIGN.CENTER)


def s03_what_is(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "What is SentinelAI?",
        "Three capabilities \u2014 built and designed on the stack we already use")

    cols = [
        ("Log RCA",                BLUE,   "Architecture-v2", [
            "Paste log or upload file",
            "AI returns: issue, root cause,",
            "  impacted service, fix steps",
            "Enriched with past incident RAG",
        ]),
        ("ELK Live Investigation", CYAN,   "ELK-Investigation-v1", [
            "Select service + severity + timeframe",
            "Live BoolQuery against Elasticsearch",
            "AI summary + suspicious log lines",
            "  \u2264 50 lines, deduped, AI-ranked",
        ]),
        ("Agentic AI  (Designed)", PURPLE, "Agentic-Flow-v1", [
            "ReAct agent loop over tools",
            "Queries ELK, Kafka, deploy history",
            "Correlates cross-system signals",
            "Files Jira, posts Slack autonomously",
        ]),
    ]

    cw = Inches(3.9)
    for i, (title, color, diag, items) in enumerate(cols):
        l = Inches(0.4) + i * (cw + Inches(0.22))
        rect(s, l, CONTENT_TOP, cw, Inches(5.0), fill=DKBLUE, line=color, lw=Pt(2))
        tb(s, title, l + Inches(0.15), CONTENT_TOP + Inches(0.12),
           cw - Inches(0.3), Inches(0.45), size=15, bold=True, color=color)
        rect(s, l + Inches(0.15), CONTENT_TOP + Inches(0.62),
             cw - Inches(0.3), Inches(0.04), fill=color)
        for j, item in enumerate(items):
            sub = item.startswith("  ")
            txt = ("    \u25e6 " + item.strip()) if sub else ("\u2022  " + item)
            tb(s, txt,
               l + Inches(0.2), CONTENT_TOP + Inches(0.75) + j * Inches(0.55),
               cw - Inches(0.4), Inches(0.52),
               size=11 if sub else 12, color=LGRAY if sub else WHITE)
        # Diagram reference tag at bottom of card
        tb(s, "diagram: SentinelAI-" + diag + ".drawio",
           l + Inches(0.12), CONTENT_TOP + Inches(4.2),
           cw - Inches(0.24), Inches(0.38),
           size=8, italic=True, color=RGBColor(0x38, 0x4C, 0x60), align=PP_ALIGN.CENTER)

    rect(s, Inches(0.4), Inches(6.5), Inches(12.5), Inches(0.55),
         fill=RGBColor(0x05, 0x18, 0x33))
    tb(s, "React  \u00b7  Spring Boot  \u00b7  FastAPI  \u00b7  PostgreSQL + pgvector  \u00b7  Elasticsearch  \u00b7  Groq / Ollama  \u00b7  Docker",
       Inches(0.6), Inches(6.57), Inches(12.2), Inches(0.44),
       size=11, color=LGRAY, align=PP_ALIGN.CENTER)


def s04_architecture(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "High-Level Architecture",
        "Three independently deployable services + Docker-compose infra")

    # Left: service layer cards
    layers = [
        ("sentinelai-ui    :3000",   BLUE,   "React SPA  \u00b7  Log RCA tab  \u00b7  ELK Investigation tab"),
        ("SentinelAI-Core  :8080",   CYAN,   "Spring Boot  \u00b7  REST APIs  \u00b7  RAG  \u00b7  ElkClient  \u00b7  pgvector"),
        ("SentinelAI-Engine  :8000", PURPLE, "FastAPI  \u00b7  Provider routing: Groq \u00b7 Ollama \u00b7 OpenAI \u00b7 HF"),
        ("Infra  (Docker Compose)",  ORANGE, "PostgreSQL+pgvector  \u00b7  Elasticsearch 7.x  \u00b7  Ollama  \u00b7  Kafka (P2)"),
    ]
    lh = Inches(1.18)
    for i, (name, color, desc) in enumerate(layers):
        t = CONTENT_TOP + i * (lh + Inches(0.1))
        rect(s, L_LEFT, t, W_LEFT, lh, fill=DKBLUE, line=color, lw=Pt(2))
        tb(s, name, L_LEFT + Inches(0.15), t + Inches(0.1),
           W_LEFT - Inches(0.3), Inches(0.42), size=13, bold=True, color=color)
        tb(s, desc, L_LEFT + Inches(0.15), t + Inches(0.56),
           W_LEFT - Inches(0.3), Inches(0.54), size=11, color=LGRAY)
        if i < len(layers) - 1:
            tb(s, "\u25bc", L_LEFT + Inches(2.4), t + lh, Inches(0.4), Inches(0.12),
               size=9, color=BLUE, align=PP_ALIGN.CENTER)

    # Right: diagram placeholder
    img_placeholder(s, L_RIGHT, CONTENT_TOP, W_RIGHT, CONTENT_H - Inches(0.15),
                    "SentinelAI-Architecture-v2.drawio")


def s05_elk(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "ELK-Based Live Investigation",
        "Select service + severity + timeframe  \u2192  AI-generated analysis in seconds")

    # Left: numbered flow steps
    steps = [
        (CYAN,   "1  Engineer selects",          "service  \u00b7  severity  \u00b7  time range in the UI tab"),
        (BLUE,   "2  POST to Core",               "POST /api/elk-investigation/search"),
        (GREEN,  "3  Core \u2192 Elasticsearch",  "BoolQuery: service + severity + date range  \u00b7  \u226450 lines"),
        (ORANGE, "4  Core \u2192 AI-Engine",       "ElkPromptBuilder numbers lines  \u2192  POST /analyze?provider=groq"),
        (YELLOW, "5  Response \u2192 UI",          "Summary  \u00b7  Root cause  \u00b7  Impacted service  \u00b7  Suspicious logs"),
    ]
    for i, (color, step, desc) in enumerate(steps):
        t = CONTENT_TOP + i * Inches(1.02)
        rect(s, L_LEFT, t, W_LEFT, Inches(0.94), fill=DKBLUE, line=color, lw=Pt(1.8))
        tb(s, step, L_LEFT + Inches(0.15), t + Inches(0.08),
           W_LEFT - Inches(0.3), Inches(0.38), size=13, bold=True, color=color)
        tb(s, desc, L_LEFT + Inches(0.15), t + Inches(0.5),
           W_LEFT - Inches(0.3), Inches(0.38), size=11, color=LGRAY)

    # Right: diagram placeholder
    img_placeholder(s, L_RIGHT, CONTENT_TOP, W_RIGHT, CONTENT_H - Inches(0.15),
                    "SentinelAI-ELK-Investigation-v1.drawio")


def s06_rca_rag(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "Log RCA  +  Incident Knowledge Base",
        "Two complementary capabilities \u2014 each makes the other more valuable")

    cw  = Inches(6.1)
    ch  = Inches(5.55)
    cw2 = Inches(6.4)

    # Card 1: Log RCA
    rect(s, L_LEFT, CONTENT_TOP, cw, ch, fill=DKBLUE, line=BLUE, lw=Pt(2))
    tb(s, "Log RCA",
       L_LEFT + Inches(0.2), CONTENT_TOP + Inches(0.1), cw - Inches(0.4), Inches(0.44),
       size=16, bold=True, color=BLUE)
    rca_items = [
        "Paste log or upload file  \u2192  POST /api/rca/analyze",
        "Optionally retrieves similar past incidents via RAG",
        "Prompt = log text + retrieved context  \u2192  AI-Engine",
        "Groq Llama 3.3 70B returns structured JSON:",
        "  issue  \u00b7  rootCause  \u00b7  impactedService",
        "  recommendedFix  \u00b7  provider  \u00b7  errors",
        "Provider swappable via config  \u2014  no code change",
    ]
    for i, item in enumerate(rca_items):
        sub = item.startswith("  ")
        txt = ("    " + item.strip()) if sub else ("\u2022  " + item)
        tb(s, txt,
           L_LEFT + Inches(0.2), CONTENT_TOP + Inches(0.65) + i * Inches(0.64),
           cw - Inches(0.4), Inches(0.58),
           size=11 if sub else 12, color=LGRAY if sub else WHITE)

    # Card 2: Knowledge Base / RAG
    l2 = L_LEFT + cw + Inches(0.25)
    rect(s, l2, CONTENT_TOP, cw2, ch, fill=DKBLUE, line=GREEN, lw=Pt(2))
    tb(s, "Incident Knowledge Base  (RAG)",
       l2 + Inches(0.2), CONTENT_TOP + Inches(0.1), cw2 - Inches(0.4), Inches(0.44),
       size=16, bold=True, color=GREEN)
    rag_items = [
        "Store incidents:  POST /api/logs/upload",
        "Core generates vector embeddings (Ollama / configurable)",
        "Stored in PostgreSQL with pgvector extension",
        "Semantic search:  POST /api/logs/search  \u2192  cosine top-K",
        "On each RCA:  top similar incidents fed to LLM as context",
        "Every stored incident improves future RCA quality",
        "Roadmap: BM25 hybrid search  \u00b7  multi-tenant  \u00b7  Qdrant",
    ]
    for i, item in enumerate(rag_items):
        tb(s, "\u2022  " + item,
           l2 + Inches(0.2), CONTENT_TOP + Inches(0.65) + i * Inches(0.64),
           cw2 - Inches(0.4), Inches(0.58), size=12, color=WHITE)


def s07_agentic(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "Agentic AI \u2014 SentinelAI as an Autonomous SRE Copilot",
        "Beyond a single LLM call: a reasoning agent that plans, uses tools, and iterates")

    # Left: ReAct loop + tool registry
    tb(s, "ReAct Agent Loop",
       L_LEFT, CONTENT_TOP, W_LEFT, Inches(0.4),
       size=14, bold=True, color=PURPLE)

    loop_steps = [
        (CYAN,   "Reason",  "Analyse context + tool results so far"),
        (ORANGE, "Act",     "Call a tool  (ELK, Kafka, deploy history, KB)"),
        (GREEN,  "Observe", "Ingest result, update working memory"),
        (YELLOW, "Repeat \u2192 Output", "Until goal met \u2192 emit RCA + runbook"),
    ]
    for i, (color, label, desc) in enumerate(loop_steps):
        t = CONTENT_TOP + Inches(0.5) + i * Inches(0.84)
        rect(s, L_LEFT, t, W_LEFT, Inches(0.76), fill=DKBLUE, line=color, lw=Pt(1.8))
        tb(s, label, L_LEFT + Inches(0.15), t + Inches(0.08),
           Inches(1.5), Inches(0.35), size=13, bold=True, color=color)
        tb(s, desc, L_LEFT + Inches(1.72), t + Inches(0.1),
           W_LEFT - Inches(1.88), Inches(0.55), size=11, color=LGRAY)

    tb(s, "Tool Registry  (what the agent can invoke)",
       L_LEFT, CONTENT_TOP + Inches(4.0), W_LEFT, Inches(0.38),
       size=12, bold=True, color=CYAN)
    tools = [
        "query_elasticsearch()", "query_kafka_lag()",
        "get_deployment_history()", "search_incident_kb()",
        "create_jira_ticket()", "post_slack_alert()",
    ]
    for i, t_name in enumerate(tools):
        col = i % 2
        row = i // 2
        tb(s, "\u2022 " + t_name,
           L_LEFT + col * Inches(2.7), CONTENT_TOP + Inches(4.45) + row * Inches(0.42),
           Inches(2.6), Inches(0.38), size=11, color=GREEN)

    # Right: diagram placeholder
    img_placeholder(s, L_RIGHT, CONTENT_TOP, W_RIGHT, CONTENT_H - Inches(0.15),
                    "SentinelAI-Agentic-Flow-v1.drawio")


def s08_orchestration(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "Agentic Orchestration \u2014 Technical Architecture",
        "Designed foundation: this is why it is not trivial")

    # Left: engineering challenges
    tb(s, "Engineering Challenges (solved in design)",
       L_LEFT, CONTENT_TOP, W_LEFT, Inches(0.4),
       size=13, bold=True, color=ORANGE)
    challenges = [
        "Context window management across many tool calls",
        "Parallel tool invocation without race conditions",
        "Prompt-injection safety from raw log content",
        "Hallucination guardrails + self-consistency checks",
        "Cost & latency budget enforcement per query",
        "Deterministic audit trail for enterprise compliance",
    ]
    for i, ch in enumerate(challenges):
        tb(s, "\u2022  " + ch,
           L_LEFT, CONTENT_TOP + Inches(0.52) + i * Inches(0.6),
           W_LEFT, Inches(0.56), size=12, color=WHITE)

    tb(s, "Two diagrams on the right show the full design.",
       L_LEFT, CONTENT_TOP + Inches(4.2), W_LEFT, Inches(0.38),
       size=11, italic=True, color=LGRAY)

    # Right: two stacked diagram placeholders
    half_h = (CONTENT_H - Inches(0.45)) / 2
    img_placeholder(s, L_RIGHT, CONTENT_TOP, W_RIGHT, half_h - Inches(0.08),
                    "SentinelAI-Agentic-Swimlane-v1.drawio")
    img_placeholder(s, L_RIGHT, CONTENT_TOP + half_h + Inches(0.12),
                    W_RIGHT, half_h - Inches(0.08),
                    "SentinelAI-Agentic-StateFlow-v1.drawio")


def s09_value_roadmap(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "Value for Atos  +  Roadmap",
        "Immediate gain today  \u00b7  Client differentiator  \u00b7  Path to enterprise product")

    # Left: 4 value cards
    values = [
        (CYAN,   "Internal Productivity",      "Faster triage, consistent RCA, less tribal knowledge"),
        (BLUE,   "Client Differentiation",     "Live AI-on-ELK demo \u2014 tangible GenAI on existing stack"),
        (GREEN,  "Capability Building",        "Applied GenAI: RAG, agents, LLMOps, vector databases"),
        (YELLOW, "Product / Accelerator Seed", "OIDC + RBAC + multi-tenant \u2192 reusable Atos SRE accelerator"),
    ]
    vh = Inches(1.2)
    for i, (color, title, desc) in enumerate(values):
        t = CONTENT_TOP + i * (vh + Inches(0.1))
        rect(s, L_LEFT, t, W_LEFT, vh, fill=DKBLUE, line=color, lw=Pt(2))
        tb(s, title, L_LEFT + Inches(0.15), t + Inches(0.1),
           W_LEFT - Inches(0.3), Inches(0.42), size=14, bold=True, color=color)
        tb(s, desc,  L_LEFT + Inches(0.15), t + Inches(0.57),
           W_LEFT - Inches(0.3), Inches(0.55), size=12, color=LGRAY)

    # Right: roadmap table
    rect(s, L_RIGHT, CONTENT_TOP, W_RIGHT, CONTENT_H - Inches(0.15),
         fill=DKBLUE, line=CYAN, lw=Pt(1.5))
    tb(s, "Roadmap",
       L_RIGHT + Inches(0.2), CONTENT_TOP + Inches(0.1),
       W_RIGHT - Inches(0.4), Inches(0.42),
       size=16, bold=True, color=CYAN)

    rect(s, L_RIGHT + Inches(0.2), CONTENT_TOP + Inches(0.58),
         W_RIGHT - Inches(0.4), Inches(0.04), fill=CYAN)

    phases = [
        ("P1  \u2014  Next 6\u201312 months", CYAN, [
            "Natural language operational search",
            "'Show payment failures after latest deploy' \u2192 AI summary",
            "Incident correlation & investigation timeline",
            "Deployment \u2192 Kafka lag \u2192 DB timeout \u2192 errors \u2192 one view",
            "Deployment regression detection",
        ]),
        ("P2  \u2014  Beyond 12 months", ORANGE, [
            "Real-time ELK/Kafka monitoring with AI anomaly detection",
            "Jira + Slack / Teams automation from AI-generated RCA",
            "Enterprise: SSO, RBAC, multi-tenant, audit logging",
            "LLMOps: feedback loop, cost tracking, A/B model testing",
        ]),
    ]

    ty = CONTENT_TOP + Inches(0.72)
    for phase_label, color, items in phases:
        tb(s, phase_label, L_RIGHT + Inches(0.2), ty,
           W_RIGHT - Inches(0.4), Inches(0.38),
           size=13, bold=True, color=color)
        ty += Inches(0.42)
        for item in items:
            tb(s, "\u2022  " + item, L_RIGHT + Inches(0.2), ty,
               W_RIGHT - Inches(0.4), Inches(0.4), size=11, color=WHITE)
            ty += Inches(0.42)
        ty += Inches(0.2)


def s10_ask_close(prs):
    s = blank_slide(prs)
    bg(s, NAVY)
    hdr(s, "Next Steps & Ask",
        "Three things from Atos stakeholders to move forward")

    # Three ask cards across the top half
    asks = [
        (CYAN,   "Identify a Pilot",
                 "One service/team to co-innovate.\nPayments, invoicing, or ops domain.\nReal ELK index with incident history."),
        (ORANGE, "Prioritise P1 Feature",
                 "Choose one to implement first:\n\u2022 Natural language search\n\u2022 Correlation timeline\n\u2022 Regression detection"),
        (GREEN,  "Support & Sponsorship",
                 "Time + ELK/Kafka env access.\nPosition as an Atos AI-SRE\nacelerator for client bids."),
    ]
    aw = Inches(3.9)
    ah = Inches(2.2)
    for i, (color, title, desc) in enumerate(asks):
        l = Inches(0.4) + i * (aw + Inches(0.22))
        rect(s, l, CONTENT_TOP, aw, ah, fill=DKBLUE, line=color, lw=Pt(2))
        tb(s, title, l + Inches(0.15), CONTENT_TOP + Inches(0.12),
           aw - Inches(0.3), Inches(0.42), size=14, bold=True, color=color)
        rect(s, l + Inches(0.15), CONTENT_TOP + Inches(0.6),
             aw - Inches(0.3), Inches(0.04), fill=color)
        tb(s, desc, l + Inches(0.15), CONTENT_TOP + Inches(0.72),
           aw - Inches(0.3), Inches(1.38), size=11, color=WHITE)

    # Thank You section
    sep_top = CONTENT_TOP + ah + Inches(0.5)
    rect(s, 0, sep_top, W, Inches(0.07), fill=BLUE)

    tb(s, "Thank You",
       Inches(1.0), sep_top + Inches(0.22), Inches(11.3), Inches(0.9),
       size=42, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
    tb(s, "Praveen Kumar Deshmukh  |  Principal Software Engineer & Architect  |  Atos",
       Inches(1.0), sep_top + Inches(1.22), Inches(11.3), Inches(0.42),
       size=14, color=LGRAY, align=PP_ALIGN.CENTER)
    tb(s, "Questions?  Let\u2019s talk about making SentinelAI real for Atos.",
       Inches(1.5), sep_top + Inches(1.74), Inches(10.3), Inches(0.42),
       size=14, italic=True, color=CYAN, align=PP_ALIGN.CENTER)


# ── Main ─────────────────────────────────────────────────────────────────────

def main():
    prs = Presentation()
    prs.slide_width  = W
    prs.slide_height = H

    slides = [
        (s01_title,         " 1 \u2013 Title"),
        (s02_problem_why,   " 2 \u2013 Problem & Why Now"),
        (s03_what_is,       " 3 \u2013 What is SentinelAI?"),
        (s04_architecture,  " 4 \u2013 Architecture              [SentinelAI-Architecture-v2.drawio]"),
        (s05_elk,           " 5 \u2013 ELK Investigation          [SentinelAI-ELK-Investigation-v1.drawio]"),
        (s06_rca_rag,       " 6 \u2013 Log RCA + Knowledge Base"),
        (s07_agentic,       " 7 \u2013 Agentic AI                 [SentinelAI-Agentic-Flow-v1.drawio]"),
        (s08_orchestration, " 8 \u2013 Orchestration              [Swimlane + StateFlow diagrams]"),
        (s09_value_roadmap, " 9 \u2013 Value + Roadmap"),
        (s10_ask_close,     "10 \u2013 Ask + Thank You"),
    ]

    print("Building SentinelAI Atos Presentation  (v2 \u2014 10 lean slides + diagram placeholders) \u2026")
    for fn, label in slides:
        fn(prs)
        print(f"  \u2713 Slide {label}")

    out = "SentinelAI_Atos_Presentation.pptx"
    prs.save(out)
    print(f"\n\u2705  Saved \u2192 {out}")
    print("\n\u2500" * 60)
    print("IMAGE PLACEHOLDERS \u2014 replace these boxes in PowerPoint / Google Slides:")
    print("  Slide  4:  SentinelAI-Architecture-v2.drawio")
    print("  Slide  5:  SentinelAI-ELK-Investigation-v1.drawio")
    print("  Slide  7:  SentinelAI-Agentic-Flow-v1.drawio")
    print("  Slide  8 top:    SentinelAI-Agentic-Swimlane-v1.drawio")
    print("  Slide  8 bottom: SentinelAI-Agentic-StateFlow-v1.drawio")
    print("\nTo export each diagram: draw.io \u2192 File \u2192 Export As \u2192 PNG (Scale: 2x)")


if __name__ == "__main__":
    main()
